"""
utils/xlsx_export.py
====================
Clean XLSX export utility for the Timesheet App.

Root causes of hidden/stale cells in exported XLSX files
---------------------------------------------------------
1. Cell-creation side-effect in the width loop: calling ws.cell(row, col)
   inside a nested for-loop creates Cell objects for every (row, col) pair
   even when there is no data. openpyxl then serialises those objects as
   <c> XML elements that external systems read as "populated".

2. Inflated worksheet dimensions: ws.max_row / ws.max_column are derived
   from every Cell ever touched, not just cells with data, so the
   <dimension ref> attribute in the XML may be e.g. A1:I1048576 even when
   only A1:I101 contains data.

3. Residual empty-cell objects: after Pandas writes the file some cells
   have value=None or value="" but still live in ws._cells. openpyxl
   serialises them as <c> elements that external systems treat as populated.

4. Incorrect style reset: the original code used
   cell._style = ws.cell(1,1)._style.__class__() which instantiates a
   generic object rather than a proper openpyxl StyleArray.

5. Pandas NaN/None serialisation: DataFrame.where(notna(), None) followed
   by to_excel can still produce empty-string cells for certain dtypes.

This module fixes all five issues.
"""

import io
import math
import openpyxl
import openpyxl.styles
import pandas as pd
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# DataFrame sanitisation
# ---------------------------------------------------------------------------

def sanitise_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Return a copy of *df* where every empty-like value is Python None.

    Treated as empty:
    - float NaN / numpy NaN (any numeric dtype)
    - pd.NaT / pd.NA
    - None
    - Strings that are empty or whitespace-only after .strip()

    Important: df.where(notna(), None) does NOT replace NaN for float/numpy
    columns because the column stays float64 and NaN is a valid float.
    We must convert to object dtype first so Python None can be stored.
    """
    df = df.copy()
    for col in df.columns:
        # Convert to object so Python None can be stored (not numpy float NaN)
        series = df[col].astype(object)

        def _clean(v):
            if v is None:
                return None
            # Catch numpy NaN, float NaN, pd.NaT, pd.NA
            try:
                if pd.isna(v):
                    return None
            except (TypeError, ValueError):
                pass
            if isinstance(v, str):
                s = v.strip()
                return s if s else None
            return v

        df[col] = series.apply(_clean)
    return df


# ---------------------------------------------------------------------------
# Column-width calculation WITHOUT cell-creation side effects
# ---------------------------------------------------------------------------

def apply_column_widths(ws, data_rows: int, data_cols: int,
                         comment_col_name: str = "Comment",
                         comment_width: int = 40):
    """Set column widths by reading only already-existing cells in ws._cells.

    The critical fix: we do NOT call ws.cell(row, col) in a nested loop.
    Calling ws.cell() creates a new Cell object as a side-effect — those
    objects are then serialised as <c> elements in the XLSX XML.
    """
    comment_col_idx = None
    col_max_len: dict = {}

    for col_idx in range(1, data_cols + 1):
        cell = ws._cells.get((1, col_idx))
        if cell and cell.value == comment_col_name:
            comment_col_idx = col_idx

    for (row, col), cell in ws._cells.items():
        if row < 1 or col < 1 or row > data_rows or col > data_cols:
            continue
        if cell.value is not None:
            length = len(str(cell.value))
            if length > col_max_len.get(col, 0):
                col_max_len[col] = length

    if comment_col_idx:
        for row_idx in range(2, data_rows + 1):
            cell = ws._cells.get((row_idx, comment_col_idx))
            if cell and cell.value is not None:
                cell.alignment = openpyxl.styles.Alignment(wrap_text=True)

    for col_idx in range(1, data_cols + 1):
        letter = get_column_letter(col_idx)
        if col_idx == comment_col_idx:
            ws.column_dimensions[letter].width = comment_width
        else:
            ws.column_dimensions[letter].width = col_max_len.get(col_idx, 8) + 2


# ---------------------------------------------------------------------------
# Deep clean — delete empty Cell objects, reset dimensions
# ---------------------------------------------------------------------------

def deep_clean_worksheet(ws, data_rows: int, data_cols: int):
    """Remove empty/stale Cell objects from ws._cells and fix dimensions."""
    keys_to_delete = []
    for (row, col), cell in list(ws._cells.items()):
        val = cell.value
        is_empty = (
            val is None
            or (isinstance(val, float) and math.isnan(val))
            or (isinstance(val, str) and not val.strip())
        )
        is_out_of_range = row > data_rows or col > data_cols
        if is_empty or is_out_of_range:
            keys_to_delete.append((row, col))

    for key in keys_to_delete:
        ws._cells.pop(key, None)

    # ws.dimensions is a read-only property in openpyxl 3.1.5 that calculates
    # the dimensions dynamically. Removing empty cells from ws._cells
    # automatically fixes the computed dimensions.

    valid_letters = {get_column_letter(c) for c in range(1, data_cols + 1)}
    for letter in list(ws.column_dimensions.keys()):
        if letter not in valid_letters:
            del ws.column_dimensions[letter]

    for r in list(ws.row_dimensions.keys()):
        if r > data_rows:
            del ws.row_dimensions[r]


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_clean_xlsx(df: pd.DataFrame,
                     sheet_name: str = "Sheet1",
                     comment_col_name: str = "Comment") -> bytes:
    """Write *df* to a structurally clean XLSX and return the bytes.

    Steps:
    1. Sanitise the DataFrame (NaN/None/whitespace -> None).
    2. Write with index=False via ExcelWriter.
    3. Post-process:
       a. Remove bold/border from header row (reading from _cells only).
       b. Calculate and set column widths without creating extra cells.
       c. Delete empty Cell objects from ws._cells.
       d. Reset worksheet dimensions to the true data bounding box.
    4. Return clean bytes.
    """
    clean_df = sanitise_dataframe(df)
    n_data_rows = len(clean_df) + 1   # +1 for header row
    n_data_cols = len(clean_df.columns)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        clean_df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        # Remove bold/border from header — read from _cells, no side effects
        for col_idx in range(1, n_data_cols + 1):
            cell = ws._cells.get((1, col_idx))
            if cell:
                cell.font = openpyxl.styles.Font(bold=False)
                cell.border = openpyxl.styles.Border()

        apply_column_widths(ws, n_data_rows, n_data_cols, comment_col_name)
        deep_clean_worksheet(ws, n_data_rows, n_data_cols)

    return buffer.getvalue()


def build_styled_summary_xlsx(df: pd.DataFrame,
                               weekday_col_names: set = None,
                               sheet_name: str = "Summary") -> bytes:
    """Write report summary DataFrame to styled XLSX matching UI specification:
    1. Header row: bold font, soft blue background (#729FCF), thin borders, center alignment for dates.
    2. Data rows:
       - Thin light gray borders (#D4D4D4) across all cells.
       - Outer perimeter border line (medium top/bottom/left/right border).
       - Weekdays (Mon-Fri, non-holiday): If hours are not shown (empty/blank/None/0), highlight in pure yellow (#FFFF00).
       - Center aligned for dates, hours, and status.
    3. Proper auto column widths and grid lines visible.
    """
    if weekday_col_names is None:
        weekday_col_names = set()

    clean_df = sanitise_dataframe(df)
    n_data_rows = len(clean_df) + 1   # +1 for header row
    n_data_cols = len(clean_df.columns)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        clean_df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        header_fill = openpyxl.styles.PatternFill(start_color="729FCF", end_color="729FCF", fill_type="solid")
        header_font = openpyxl.styles.Font(name="Calibri", size=11, bold=True, color="000000")
        yellow_fill = openpyxl.styles.PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        for r_idx in range(1, n_data_rows + 1):
            for c_idx in range(1, n_data_cols + 1):
                cell = ws.cell(row=r_idx, column=c_idx)

                top_style = "medium" if r_idx == 1 else "thin"
                bottom_style = "medium" if r_idx == n_data_rows else "thin"
                left_style = "medium" if c_idx == 1 else "thin"
                right_style = "medium" if c_idx == n_data_cols else "thin"

                top_color = "000000" if r_idx == 1 else "D4D4D4"
                bottom_color = "000000" if r_idx == n_data_rows else "D4D4D4"
                left_color = "000000" if c_idx == 1 else "D4D4D4"
                right_color = "000000" if c_idx == n_data_cols else "D4D4D4"

                cell.border = openpyxl.styles.Border(
                    top=openpyxl.styles.Side(style=top_style, color=top_color),
                    bottom=openpyxl.styles.Side(style=bottom_style, color=bottom_color),
                    left=openpyxl.styles.Side(style=left_style, color=left_color),
                    right=openpyxl.styles.Side(style=right_style, color=right_color)
                )

                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = openpyxl.styles.Alignment(
                        horizontal="center" if c_idx > 2 else "left",
                        vertical="center"
                    )
                else:
                    col_name = clean_df.columns[c_idx - 1]
                    if c_idx > 2:
                        cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

                    # Highlight missing weekday hours in yellow
                    if col_name in weekday_col_names:
                        v = cell.value
                        if v is None or str(v).strip() in ("", "none", "nan") or v == 0:
                            cell.fill = yellow_fill

        apply_column_widths(ws, n_data_rows, n_data_cols)
        ws.views.sheetView[0].showGridLines = True

    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Validation helper (used by tests)
# ---------------------------------------------------------------------------

def validate_xlsx(xlsx_bytes: bytes,
                  expected_rows: int,
                  expected_cols: int,
                  sheet_name: str = "Sheet1") -> dict:
    """Programmatically validate a generated XLSX for cleanliness.

    Parameters
    ----------
    xlsx_bytes     : XLSX file bytes to validate.
    expected_rows  : total expected rows including the header row.
    expected_cols  : expected number of columns.
    sheet_name     : sheet to inspect.

    Returns dict with: ok, issues, unexpected_cells, used_rows, used_cols.
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[sheet_name]

    issues = []
    unexpected = []
    max_row_found = 0
    max_col_found = 0

    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            has_value = (
                val is not None
                and not (isinstance(val, float) and math.isnan(val))
                and not (isinstance(val, str) and not val.strip())
            )
            if has_value:
                if cell.row > max_row_found:
                    max_row_found = cell.row
                if cell.column > max_col_found:
                    max_col_found = cell.column
                if cell.row > expected_rows or cell.column > expected_cols:
                    unexpected.append("{coord}={val!r}".format(
                        coord=cell.coordinate, val=val))

    if unexpected:
        issues.append(
            "{n} unexpected cell(s) outside A1:{col}{row}: {cells}".format(
                n=len(unexpected),
                col=get_column_letter(expected_cols),
                row=expected_rows,
                cells=", ".join(unexpected[:10]),
            )
        )

    return {
        "ok": len(issues) == 0,
        "issues": issues,
        "unexpected_cells": unexpected,
        "used_rows": max_row_found,
        "used_cols": max_col_found,
    }

