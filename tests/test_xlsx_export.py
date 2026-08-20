"""
tests/test_xlsx_export.py
=========================
Regression tests for the clean XLSX export utility.

Tests cover all six scenarios from the requirements:
  Test 1 - Normal export (no unexpected cells outside A1:E11)
  Test 2 - NaN values produce truly empty cells
  Test 3 - Empty-string values produce truly empty cells
  Test 4 - Whitespace-only values produce truly empty cells
  Test 5 - No cells outside the data range are populated
  Test 6 - validate_xlsx confirms structural cleanliness
"""

import io
import math
import unittest
import pandas as pd
import openpyxl

import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from utils.xlsx_export import build_clean_xlsx, sanitise_dataframe, validate_xlsx


def _load_ws(xlsx_bytes, sheet_name="Sheet1"):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    return wb[sheet_name]


def _all_cells_with_value(ws):
    result = []
    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            if val is None:
                continue
            if isinstance(val, float) and math.isnan(val):
                continue
            if isinstance(val, str) and not val.strip():
                continue
            result.append((cell.coordinate, val))
    return result


class TestNormalExport(unittest.TestCase):
    """Test 1 - Given 10 rows x 5 cols, XLSX must not exceed A1:E11."""

    def setUp(self):
        self.df = pd.DataFrame({
            'Col_A': [f'row{i}' for i in range(10)],
            'Col_B': list(range(10)),
            'Col_C': [f'val{i}' for i in range(10)],
            'Col_D': [float(i) * 1.5 for i in range(10)],
            'Col_E': [i % 3 for i in range(10)],
        })
        self.xlsx = build_clean_xlsx(self.df, sheet_name='Sheet1')

    def test_no_cells_outside_range(self):
        report = validate_xlsx(self.xlsx, expected_rows=11, expected_cols=5)
        self.assertTrue(report['ok'], msg=str(report['issues']))

    def test_used_rows_match(self):
        report = validate_xlsx(self.xlsx, expected_rows=11, expected_cols=5)
        self.assertEqual(report['used_rows'], 11)

    def test_used_cols_match(self):
        report = validate_xlsx(self.xlsx, expected_rows=11, expected_cols=5)
        self.assertEqual(report['used_cols'], 5)

    def test_worksheet_dimensions_correct(self):
        ws = _load_ws(self.xlsx)
        self.assertEqual(ws.dimensions, 'A1:E11')


class TestNaNValues(unittest.TestCase):
    """Test 2 - NaN values produce genuinely empty cells."""

    def setUp(self):
        self.df = pd.DataFrame({
            'Name':  ['John', 'Jane', 'Bob'],
            'Hours': [8.0,    float('nan'), 7.0],
        })
        self.xlsx = build_clean_xlsx(self.df, sheet_name='Sheet1')

    def test_nan_cell_is_none(self):
        ws = _load_ws(self.xlsx)
        jane_hours = ws.cell(row=3, column=2).value
        self.assertIsNone(jane_hours)

    def test_validation_passes(self):
        report = validate_xlsx(self.xlsx, expected_rows=4, expected_cols=2)
        self.assertTrue(report['ok'], msg=str(report['issues']))


class TestEmptyStringValues(unittest.TestCase):
    """Test 3 - Empty-string values produce genuinely empty cells."""

    def setUp(self):
        self.df = pd.DataFrame({
            'Name':  ['John', 'Jane', 'Bob'],
            'Hours': ['8',    '',     '7'],
        })
        self.xlsx = build_clean_xlsx(self.df, sheet_name='Sheet1')

    def test_empty_string_cell_is_none(self):
        ws = _load_ws(self.xlsx)
        jane_hours = ws.cell(row=3, column=2).value
        self.assertIsNone(jane_hours)

    def test_validation_passes(self):
        report = validate_xlsx(self.xlsx, expected_rows=4, expected_cols=2)
        self.assertTrue(report['ok'], msg=str(report['issues']))


class TestWhitespaceValues(unittest.TestCase):
    """Test 4 - Whitespace-only cells are cleared."""

    def setUp(self):
        self.df = pd.DataFrame({'Name': ['Jane'], 'Value': ['   ']})
        self.xlsx = build_clean_xlsx(self.df, sheet_name='Sheet1')

    def test_whitespace_cell_is_none(self):
        ws = _load_ws(self.xlsx)
        val = ws.cell(row=2, column=2).value
        self.assertIsNone(val)

    def test_validation_passes(self):
        report = validate_xlsx(self.xlsx, expected_rows=2, expected_cols=2)
        self.assertTrue(report['ok'], msg=str(report['issues']))


class TestNoExtraCells(unittest.TestCase):
    """Test 5 - No cells outside the data range are populated."""

    def setUp(self):
        self.df = pd.DataFrame({'A': [1, 2], 'B': ['x', 'y']})
        self.xlsx = build_clean_xlsx(self.df, sheet_name='Sheet1')

    def test_no_extra_columns(self):
        ws = _load_ws(self.xlsx)
        for row in ws.iter_rows():
            for cell in row:
                if cell.column > 2:
                    self.assertIsNone(cell.value,
                        msg=f"Unexpected value in {cell.coordinate}")

    def test_no_extra_rows(self):
        ws = _load_ws(self.xlsx)
        for row in ws.iter_rows():
            for cell in row:
                if cell.row > 3:
                    self.assertIsNone(cell.value,
                        msg=f"Unexpected value in {cell.coordinate}")

    def test_z100_is_empty(self):
        ws = _load_ws(self.xlsx)
        cell = ws._cells.get((100, 26))
        self.assertIsNone(cell)

    def test_dimensions_tight(self):
        ws = _load_ws(self.xlsx)
        self.assertEqual(ws.dimensions, 'A1:B3')


class TestValidateXlsx(unittest.TestCase):
    """Test 6 - validate_xlsx confirms structural cleanliness."""

    def test_clean_file_passes(self):
        df = pd.DataFrame({'X': [1, 2, 3], 'Y': ['a', 'b', 'c']})
        xlsx = build_clean_xlsx(df)
        report = validate_xlsx(xlsx, expected_rows=4, expected_cols=2)
        self.assertTrue(report['ok'])

    def test_dirty_file_fails(self):
        df = pd.DataFrame({'X': [1], 'Y': ['a']})
        xlsx_bytes = build_clean_xlsx(df)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        ws['Z100'] = 'stale'
        buf = io.BytesIO()
        wb.save(buf)
        dirty = buf.getvalue()
        report = validate_xlsx(dirty, expected_rows=2, expected_cols=2)
        self.assertFalse(report['ok'])
        self.assertTrue(any('Z100' in c for c in report['unexpected_cells']))


class TestSanitiseDataframe(unittest.TestCase):

    def test_nan_becomes_none(self):
        df = pd.DataFrame({'v': [float('nan')]})
        r = sanitise_dataframe(df)
        self.assertIsNone(r.iloc[0, 0])

    def test_empty_string_becomes_none(self):
        df = pd.DataFrame({'v': ['']})
        r = sanitise_dataframe(df)
        self.assertIsNone(r.iloc[0, 0])

    def test_whitespace_becomes_none(self):
        df = pd.DataFrame({'v': ['   ']})
        r = sanitise_dataframe(df)
        self.assertIsNone(r.iloc[0, 0])

    def test_valid_string_preserved(self):
        df = pd.DataFrame({'v': ['hello']})
        r = sanitise_dataframe(df)
        self.assertEqual(r.iloc[0, 0], 'hello')

    def test_stripped_string_preserved(self):
        df = pd.DataFrame({'v': ['  hello  ']})
        r = sanitise_dataframe(df)
        self.assertEqual(r.iloc[0, 0], 'hello')

    def test_original_not_mutated(self):
        df = pd.DataFrame({'v': ['  ']})
        _ = sanitise_dataframe(df)
        self.assertEqual(df.iloc[0, 0], '  ')


if __name__ == '__main__':
    unittest.main()
